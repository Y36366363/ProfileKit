import os
import unittest
from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch

from fastapi.testclient import TestClient
from pypdf import PdfReader
from docx import Document

from profilekit.models import AgentTurn, ProfileRecord, ProfileSession, WorkflowStage
from profilekit.web import SessionStore, app
import profilekit.web as web_module


class WebAppTests(unittest.TestCase):
    def setUp(self):
        self.directory = TemporaryDirectory()
        self.original_store = web_module.store
        web_module.store = SessionStore(Path(self.directory.name) / "session.json")
        self.client = TestClient(app)

    def tearDown(self):
        web_module.store = self.original_store
        self.directory.cleanup()

    def test_home_page_contains_primary_workspace(self):
        response = self.client.get("/")
        self.assertEqual(response.status_code, 200)
        self.assertIn("ProfileKit", response.text)
        self.assertIn("Personal Profile Record", response.text)
        self.assertIn("From sources to one clear page", response.text)

    def test_public_state_never_exposes_api_key(self):
        response = self.client.get("/api/session")
        self.assertEqual(response.status_code, 200)
        body = response.text
        openai_key = os.environ.get("OPENAI_API_KEY")
        if openai_key:
            self.assertNotIn(openai_key, body)
        self.assertNotIn("DEEPSEEK_API_KEY", body)
        self.assertNotIn("GEMINI_API_KEY", body)
        self.assertGreaterEqual(len(response.json()["runtime"]["models"]), 5)

    def test_model_can_be_changed_for_the_local_session(self):
        with patch.dict(os.environ, {"OPENAI_API_KEY": "test-key"}, clear=False):
            response = self.client.patch(
                "/api/model",
                json={"provider": "openai", "model": "gpt-6-luna"},
            )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()["model_provider"], "openai")
        self.assertEqual(response.json()["model_name"], "gpt-6-luna")

    def test_unknown_model_is_rejected(self):
        response = self.client.patch(
            "/api/model",
            json={"provider": "openai", "model": "made-up-model"},
        )
        self.assertEqual(response.status_code, 400)

    def test_demo_loads_fictional_profile_record(self):
        response = self.client.post("/api/demo")
        self.assertEqual(response.status_code, 200)
        payload = response.json()
        self.assertEqual(payload["stage"], "personal_profile_record")
        self.assertGreater(len(payload["record"]["items"]), 0)
        self.assertTrue(all("demo" in (item.get("source") or "").lower() for item in payload["record"]["items"]))
        self.assertEqual(payload["presentation"]["title"], "Lin Chen")
        self.assertNotIn("lin.chen@example.edu", str(payload["presentation"]))

    def test_theme_selection_updates_preview_and_persists(self):
        self.client.post("/api/demo")
        response = self.client.patch("/api/theme", json={"theme": "modern"})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()["profile_theme"], "modern")
        self.assertEqual(response.json()["presentation"]["theme"], "modern")
        saved = ProfileSession.model_validate_json(web_module.store.path.read_text())
        self.assertEqual(saved.profile_theme, "modern")

    def test_pdf_export_is_one_page_and_omits_unconfirmed_email(self):
        self.client.post("/api/demo")
        response = self.client.get("/api/export/pdf")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.headers["content-type"], "application/pdf")
        reader = PdfReader(BytesIO(response.content))
        self.assertEqual(len(reader.pages), 1)
        text = reader.pages[0].extract_text()
        self.assertIn("Lin Chen", text)
        self.assertNotIn("lin.chen@example.edu", text)

    def test_empty_profile_cannot_export_pdf(self):
        response = self.client.get("/api/export/pdf")
        self.assertEqual(response.status_code, 400)

    def test_item_decision_is_saved_with_private_permissions(self):
        self.client.post("/api/demo")
        response = self.client.patch("/api/items/0", json={"decision": "include"})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()["record"]["items"][0]["user_decision"], "include")
        self.assertEqual(web_module.store.path.stat().st_mode & 0o777, 0o600)

    def test_reset_returns_to_intake(self):
        web_module.store.session = ProfileSession()
        self.client.post("/api/demo")
        response = self.client.post("/api/reset")
        self.assertEqual(response.json()["stage"], "intake")

    def test_text_upload_is_queued_for_source_review(self):
        response = self.client.post(
            "/api/upload",
            files={"files": ("notes.txt", b"Supported interview coding.", "text/plain")},
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()["accepted"], ["notes.txt"])
        self.assertEqual(response.json()["session"]["pending_source_count"], 1)

    def test_docx_resume_upload_is_queued_for_agent_review(self):
        buffer = BytesIO()
        document = Document()
        document.add_heading("Jordan Rivera", level=1)
        document.add_paragraph("Graduate student in information science")
        document.add_paragraph("Project: evaluated a library search prototype")
        document.save(buffer)
        response = self.client.post(
            "/api/upload",
            files={
                "files": (
                    "resume.docx",
                    buffer.getvalue(),
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                )
            },
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()["accepted"], ["resume.docx"])
        self.assertEqual(response.json()["session"]["pending_source_count"], 1)

    def test_uploaded_source_is_hidden_from_the_visible_transcript(self):
        secret_source = "Private source phone 585-483-4876 and test@example.edu"
        self.client.post(
            "/api/upload",
            files={"files": ("private.txt", secret_source, "text/plain")},
        )

        calls = []

        def fake_turn(session, user_message, provider, model, transcript_user_message=None):
            calls.append(user_message)
            self.assertIn(secret_source, user_message)
            self.assertNotIn(secret_source, transcript_user_message)
            turn = AgentTurn(
                assistant_message="Contact details: 585-483-4876 and test@example.edu",
                proposed_stage=WorkflowStage.SOURCE_REVIEW,
                record=ProfileRecord(),
            )
            from profilekit.workflow import apply_turn

            updated = apply_turn(session, user_message, turn, transcript_user_message)
            return updated, turn.assistant_message

        with patch("profilekit.web.run_turn", side_effect=fake_turn):
            response = self.client.post("/api/chat", json={"message": "Review my file."})
            follow_up = self.client.post("/api/chat", json={"message": "Continue the review."})

        self.assertEqual(response.status_code, 200)
        self.assertEqual(follow_up.status_code, 200)
        self.assertEqual(len(calls), 2)
        visible = str(follow_up.json()["session"]["transcript"])
        self.assertNotIn(secret_source, visible)
        self.assertNotIn("585-483-4876", visible)
        self.assertNotIn("test@example.edu", visible)
        self.assertIn("privately reviewed 1 uploaded file", visible)
        self.assertIn("[phone withheld]", visible)
        self.assertIn("[email withheld]", visible)

    def test_default_config_upload_applies_profile_without_model_call(self):
        self.client.post("/api/demo")
        config = {
            "profile_metadata": {"audience": "Course faculty", "purpose": "Project showcase"},
            "profile": {
                "name": "Jordan Rivera",
                "role": "Information science student",
                "introduction": "Studies accessible information systems.",
            },
            "items": [],
            "design": {
                "theme": "minimal",
                "accent_color": "#336699",
                "font_style": "sans",
                "layout_density": "airy",
            },
        }
        response = self.client.post(
            "/api/upload",
            files={"files": ("default_config.json", __import__("json").dumps(config), "application/json")},
        )
        self.assertEqual(response.status_code, 200)
        payload = response.json()
        self.assertEqual(payload["configured"], ["default_config.json"])
        self.assertEqual(payload["session"]["presentation"]["title"], "Jordan Rivera")
        self.assertEqual(payload["session"]["profile_theme"], "minimal")
        self.assertEqual(payload["session"]["accent_color"], "#336699")
        self.assertEqual(payload["session"]["pending_source_count"], 0)
        self.assertNotIn("lin.chen@example.edu", str(payload["session"]))

    def test_preferences_form_updates_content_and_design(self):
        response = self.client.put(
            "/api/preferences",
            json={
                "name": "Avery Kim",
                "role": "UX research student",
                "introduction": "Explores understandable public services.",
                "audience": "Faculty reviewers",
                "occasion": "Course showcase",
                "purpose": "Present one research project",
                "tone": "Clear and modest",
                "visual_preferences": "Warm accent and spacious layout",
                "privacy_restrictions": "No phone number",
                "theme": "modern",
                "accent_color": "#A34F2A",
                "font_style": "sans",
                "layout_density": "airy",
            },
        )
        self.assertEqual(response.status_code, 200)
        payload = response.json()
        self.assertEqual(payload["presentation"]["title"], "Avery Kim")
        self.assertEqual(payload["presentation"]["accent_color"], "#A34F2A")
        self.assertEqual(payload["record"]["profile_metadata"]["audience"], "Faculty reviewers")


if __name__ == "__main__":
    unittest.main()
