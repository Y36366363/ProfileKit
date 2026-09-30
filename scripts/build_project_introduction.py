from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "output" / "docs" / "ProfileKit_Project_Introduction.docx"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    tc_pr.append(shading)


def set_cell_borders(cell, color: str = "D9D9D9") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:color"), color)
        borders.append(element)


def set_cell_margins(cell, top: int = 70, start: int = 105, bottom: int = 70, end: int = 105) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for key, value in {"top": top, "start": start, "bottom": bottom, "end": end}.items():
        node = OxmlElement(f"w:{key}")
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")
        margins.append(node)


def set_run_font(run, name: str = "Aptos", size: float | None = None, bold: bool | None = None) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    set_run_font(run, size=9)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = "PAGE"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, end])


def remove_paragraph_border(paragraph_or_style) -> None:
    p_pr = paragraph_or_style._element.get_or_add_pPr()
    border = p_pr.find(qn("w:pBdr"))
    if border is not None:
        p_pr.remove(border)


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    paragraph = doc.add_heading(text, level=level)
    paragraph.paragraph_format.keep_with_next = True


def add_body(doc: Document, text: str, *, bold_lead: str | None = None) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.keep_together = True
    if bold_lead and text.startswith(bold_lead):
        lead = paragraph.add_run(bold_lead)
        lead.bold = True
        paragraph.add_run(text[len(bold_lead):])
    else:
        paragraph.add_run(text)


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        paragraph = doc.add_paragraph(style="List Bullet")
        paragraph.add_run(item)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[float]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.autofit = False
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.width = Inches(widths[index])
        set_cell_shading(cell, "1F3A5F")
        set_cell_borders(cell)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run = paragraph.add_run(header)
        run.font.color.rgb = RGBColor(255, 255, 255)
        set_run_font(run, size=8.5, bold=True)
    for row_index, values in enumerate(rows):
        row = table.add_row()
        cant_split = OxmlElement("w:cantSplit")
        row._tr.get_or_add_trPr().append(cant_split)
        cells = row.cells
        for index, value in enumerate(values):
            cell = cells[index]
            cell.width = Inches(widths[index])
            set_cell_borders(cell)
            set_cell_margins(cell)
            if row_index % 2:
                set_cell_shading(cell, "F2F6FA")
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            paragraph = cell.paragraphs[0]
            run = paragraph.add_run(value)
            set_run_font(run, size=8.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def build_document() -> None:
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.72)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    normal.font.size = Pt(9.5)
    normal.paragraph_format.space_after = Pt(3)
    normal.paragraph_format.line_spacing = 1.03
    for style_name, size in (("Title", 24), ("Heading 1", 15), ("Heading 2", 12)):
        style = styles[style_name]
        style.font.name = "Aptos Display" if style_name != "Normal" else "Aptos"
        style._element.rPr.rFonts.set(qn("w:ascii"), style.font.name)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), style.font.name)
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.bold = style_name != "Title"
    styles["Title"].font.bold = True
    styles["Title"].paragraph_format.space_after = Pt(6)
    remove_paragraph_border(styles["Title"])
    styles["Heading 1"].paragraph_format.space_before = Pt(8)
    styles["Heading 1"].paragraph_format.space_after = Pt(4)
    styles["Heading 2"].paragraph_format.space_before = Pt(11)
    styles["Heading 2"].paragraph_format.space_after = Pt(5)

    header = section.header.paragraphs[0]
    header.text = "PROFILEKIT PROJECT DOCUMENTATION"
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    header_run = header.runs[0]
    set_run_font(header_run, size=8.5, bold=True)
    header_run.font.color.rgb = RGBColor(0, 0, 0)
    add_page_number(section.footer.paragraphs[0])

    title = doc.add_paragraph(style="Title")
    title.add_run("ProfileKit Project Introduction")
    remove_paragraph_border(title)
    subtitle = doc.add_paragraph()
    subtitle_run = subtitle.add_run("A privacy-conscious local agent for editable single-page personal profiles")
    set_run_font(subtitle_run, size=11.5)
    subtitle_run.font.color.rgb = RGBColor(0, 0, 0)
    meta = doc.add_paragraph()
    meta.add_run("Purpose, core functions, and classroom demonstration appendix\n")
    meta.add_run("Current status: locally runnable and tested")
    for run in meta.runs:
        set_run_font(run, size=8.5)
        run.font.color.rgb = RGBColor(0, 0, 0)

    add_heading(doc, "Purpose and Motivation")
    add_body(
        doc,
        "My goal is to build a practical agent that helps students, researchers, and early-career professionals turn their own approved information into a concise, attractive profile. A conventional resume is often too long for a class introduction, conference page, or project showcase, while an unrestricted AI tool may invent claims or expose details that are unsuitable for a public audience. ProfileKit therefore combines language-model assistance with explicit review, privacy, and design controls.",
    )
    add_body(
        doc,
        "The current version runs as a local web application. It keeps the API keys on the server, stores the working session locally, distinguishes source facts from suggested wording, and lets the user decide what appears in the final one-page PDF. It is intended as a classroom-ready prototype, not an automatic publishing or identity-verification service.",
    )

    add_heading(doc, "How the Agent Works")
    add_body(
        doc,
        "Input documents and user preferences are reviewed by the model and converted into structured profile items. A deterministic controller enforces approval stages; the user resolves sensitive or uncertain content, then the live preview and PDF use the same approved record and design settings.",
    )

    add_heading(doc, "Core Functions")
    add_table(
        doc,
        ["Function", "What the user can do"],
        [
            ["Flexible input", "Upload a CV/DOCX/PDF, describe the goal in conversation, or import default_config.json for a repeatable no-model setup."],
            ["Review and privacy", "Separate facts from suggested wording; include, exclude, revise, or restrict items, with extra safeguards for sensitive details."],
            ["Editable design", "Control audience, purpose, tone, introduction, theme, accent color, typography, and density in a live preview."],
            ["Model and output", "Use low-cost DeepSeek Flash by default (or an available OpenAI model) and export the approved result as a one-page PDF."],
        ],
        [1.35, 5.45],
    )

    doc.add_page_break()
    add_heading(doc, "Appendix: Classroom Demonstration and Validation")
    add_body(doc, "Suggested five-minute demonstration:")
    add_table(
        doc,
        ["Time", "Action", "What to explain"],
        [
            ["0:00", "Open the local site; choose DeepSeek Flash.", "English-first interface; keys stay server-side."],
            ["0:40", "Upload a sample CV or load the fictional demo; request an audience-specific profile.", "Sources become evidence and structured items, not automatic permission to publish every detail."],
            ["2:00", "Compare Record and Live Preview; exclude an item and change visual settings.", "Private review data is separate from public output, and the user controls both content and design."],
            ["4:10", "Export and open the PDF.", "The artifact is one page and reflects the approved visible state."],
        ],
        [0.65, 2.0, 4.15],
    )

    add_heading(doc, "Current Validation Status")
    add_table(
        doc,
        ["Check", "Verified result"],
        [
            ["Automated behavior", "31 tests cover workflow gates, privacy, models, configuration, uploads, customization, and one-page export."],
            ["Local web app", "127.0.0.1:8765 supports the complete browser demonstration."],
            ["Real CV workflow", "A DOCX CV can be parsed, reviewed by DeepSeek Flash, converted to a profile, and exported locally."],
            ["Privacy and fallback", ".env, sessions, and personal outputs are ignored by Git; the fictional demo and default_config.json work without a model call."],
        ],
        [1.45, 5.35],
    )

    add_heading(doc, "Known Limitations")
    add_bullets(
        doc,
        [
            "Gemini credentials can be tested, but Gemini is not yet a conversation option in the web selector.",
            "Image interpretation, approved portrait placement, and drag-and-drop section reordering remain future work.",
            "The exporter prioritizes a single page and may omit lower-priority overflow; users must review the final PDF.",
            "Formal accessibility, institutional policy, and factual accuracy still require human review.",
        ],
    )

    add_heading(doc, "Preclass Checklist")
    add_body(
        doc,
        "Run python -m profilekit.preflight, open http://127.0.0.1:8765, load the fictional demo, test Customize, and export one PDF. Keep default_config.json available as the no-model fallback, and never display the .env file or API keys during class.",
    )

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()
